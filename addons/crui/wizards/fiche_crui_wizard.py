from odoo import models, fields
from odoo.exceptions import UserError
import base64, copy, os, logging
from io import BytesIO
from collections import defaultdict
from lxml import etree

_logger = logging.getLogger(__name__)

NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'

ORDRE_PREFECTURES = [
    'FES', 'MOULAY YACOUB', 'SEFROU', 'BOULEMANE',
    'TAOUNATE', 'TAZA', 'EL HAJEB', 'IFRANE', 'MEKNES',
]
PREFECTURES_SET = {'FES', 'MEKNES'}


class FicheCruiWizard(models.TransientModel):
    _name = 'fiche.crui.wizard'
    _description = "Wizard pour générer la fiche CRUI PPTX"

    date_passage = fields.Date(
        string="Date de passage CRUI",
        required=True,
        help="Sélectionnez la date de passage pour générer les fiches CRUI"
    )
    file_data = fields.Binary(string="Fichier généré", readonly=True)
    file_name = fields.Char(string="Nom du fichier", readonly=True)

    # ------------------------------------------------------------------
    # Action principale
    # ------------------------------------------------------------------

    def action_generate_fiche_crui(self):
        self.ensure_one()
        if not self.date_passage:
            raise UserError("Veuillez sélectionner une date de passage.")

        dossiers = self.env['dossier'].sudo().search([
            ('date_passage', '>=', fields.Datetime.to_string(
                fields.Datetime.from_string(str(self.date_passage) + ' 00:00:00')
            )),
            ('date_passage', '<=', fields.Datetime.to_string(
                fields.Datetime.from_string(str(self.date_passage) + ' 23:59:59')
            )),
        ])

        if not dossiers:
            raise UserError(
                f"Aucun dossier trouvé pour la date {self.date_passage.strftime('%d/%m/%Y')}."
            )

        dossier_list = self._prepare_dossier_data(dossiers)
        grouped, sorted_prefectures = self._group_and_sort(dossier_list)

        try:
            from pptx import Presentation
        except ImportError:
            raise UserError(
                "La bibliothèque python-pptx n'est pas installée. "
                "Veuillez reconstruire l'image Docker avec python-pptx."
            )

        template_path = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            'doc', 'model_de_fiche crui.pptx'
        )
        prs = Presentation(template_path)

        self._fill_slide1(prs.slides[0], self.date_passage, grouped, sorted_prefectures)

        template_slide2 = prs.slides[1]
        all_dossiers_ordered = []
        for pref in sorted_prefectures:
            all_dossiers_ordered.extend(grouped[pref])

        # Créer d'abord tous les slides vides à partir du template non modifié,
        # puis les remplir — évite que les images du projet N se copient sur N+1
        slides_to_fill = [template_slide2]
        for _ in all_dossiers_ordered[1:]:
            slides_to_fill.append(self._duplicate_slide(prs, template_slide2))

        for idx, (slide, data) in enumerate(zip(slides_to_fill, all_dossiers_ordered)):
            self._fill_slide2(slide, idx + 1, data)

        buf = BytesIO()
        prs.save(buf)
        buf.seek(0)
        file_data = base64.b64encode(buf.read())

        date_str = self.date_passage.strftime('%d-%m-%Y')
        file_name = f"Fiche_CRUI_{date_str}.pptx"

        self.write({'file_data': file_data, 'file_name': file_name})

        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content/fiche.crui.wizard/{self.id}/file_data'
                   f'?download=true&filename={file_name}',
            'target': 'self',
        }

    # ------------------------------------------------------------------
    # Préparation des données
    # ------------------------------------------------------------------

    def _prepare_dossier_data(self, dossiers):
        result = []
        for dossier in dossiers:
            projet = dossier.projet_id

            if projet:
                prefecture_name = (
                    projet.prefecture_id.name.upper() if projet.prefecture_id else ""
                )

                # Province / Commune
                if projet.prefecture_id:
                    prefix = "Préfecture" if prefecture_name in PREFECTURES_SET else "Province"
                    pref_label = f"{prefix} de {projet.prefecture_id.name}"
                else:
                    pref_label = ""
                commune_label = (
                    f"Commune de {projet.commune_id.name}" if projet.commune_id else ""
                )
                location_str = (
                    f"{pref_label} / {commune_label}" if pref_label and commune_label
                    else pref_label or commune_label
                )

                # Foncier
                foncier_parts = list(projet.fonciers.mapped('name')) if projet.fonciers else []
                if projet.foncier_id and projet.foncier_id.name not in foncier_parts:
                    foncier_parts.append(projet.foncier_id.name)
                if projet.reference_fonciere:
                    foncier_parts.append(projet.reference_fonciere)
                foncier_str = ' / '.join(foncier_parts)

                mt_invest_mdh = (
                    projet.montant_investissement / 1_000_000
                    if projet.montant_investissement else 0
                )

                # Historique
                if projet.historique_projet:
                    historique_str = projet.historique_projet
                else:
                    dossiers_ant = self.env['dossier'].search([
                        ('projet_id', '=', projet.id),
                        ('date_passage', '<', dossier.date_passage),
                        ('id', '!=', dossier.id),
                    ], order='date_passage desc')
                    passages = []
                    for d in dossiers_ant:
                        if d.date_passage:
                            passages.append(
                                f"CRUI du {d.date_passage.strftime('%d/%m/%Y')} : "
                                f"{d.decision_crui or 'Décision non renseignée'}"
                            )
                    historique_str = " | ".join(passages) if passages else ""

                petitionnaire_str = projet.raison_social or projet.petitionnaire or ""
                projet_name = projet.name or ""
            else:
                # Dossier sans projet associé — fiche vide avec référence dossier
                prefecture_name = ""
                location_str = ""
                foncier_str = ""
                mt_invest_mdh = 0
                historique_str = ""
                petitionnaire_str = ""
                projet_name = dossier.name or ""

            # Actes (sur le dossier, pas le projet)
            actes_str = (
                ", ".join(dossier.actes_ids.mapped('name'))
                if dossier.actes_ids else ""
            )

            result.append({
                'dossier': dossier,
                'projet': projet,
                'prefecture': prefecture_name or "Hors Plateforme",
                'location_str': location_str,
                'foncier_str': foncier_str,
                'actes_str': actes_str,
                'mt_invest_mdh': mt_invest_mdh,
                'historique_str': historique_str,
                'petitionnaire_str': petitionnaire_str,
                'projet_name': projet_name,
            })
        return result

    def _group_and_sort(self, dossier_list):
        grouped = defaultdict(list)
        for data in dossier_list:
            grouped[data['prefecture']].append(data)

        sorted_prefectures = []
        for pref in ORDRE_PREFECTURES:
            if pref in grouped:
                sorted_prefectures.append(pref)
        for pref in grouped:
            if pref not in sorted_prefectures:
                sorted_prefectures.append(pref)

        return grouped, sorted_prefectures

    # ------------------------------------------------------------------
    # Remplissage Slide 1
    # ------------------------------------------------------------------

    def _fill_slide1(self, slide, date_obj, grouped, sorted_prefectures):
        for shape in slide.shapes:
            if shape.name == 'ZoneTexte 8':
                # Para index 1, run index 1 contient "/2026" → remplacer par la date complète
                paras = shape.text_frame.paragraphs
                if len(paras) > 1 and len(paras[1].runs) > 1:
                    paras[1].runs[1].text = date_obj.strftime('%d/%m/%Y')
            elif shape.name == 'Tableau 10':
                self._rebuild_summary_table(shape, grouped, sorted_prefectures)

    def _rebuild_summary_table(self, shape, grouped, sorted_prefectures):
        tbl_elem = shape.table._tbl
        rows = tbl_elem.findall(f'{{{NS_A}}}tr')

        # Sauvegarder les lignes-types avant suppression
        tmpl_multi_first = copy.deepcopy(rows[1])   # préfecture multi-actes (rowSpan)
        tmpl_multi_cont = copy.deepcopy(rows[2])    # continuation (vMerge)
        tmpl_single = copy.deepcopy(rows[7])        # mono-acte (sans span)
        tmpl_total = copy.deepcopy(rows[-1])        # ligne Total

        # Supprimer toutes les lignes de données (garder uniquement header row 0)
        for row in rows[1:]:
            tbl_elem.remove(row)

        total_count = 0
        for pref_name in sorted_prefectures:
            pref_data = grouped[pref_name]
            pref_count = len(pref_data)
            total_count += pref_count

            # Collecter les actes uniques de cette préfecture
            actes_set = []
            seen = set()
            for d in pref_data:
                for acte_name in d['dossier'].actes_ids.mapped('name'):
                    if acte_name not in seen:
                        actes_set.append(acte_name)
                        seen.add(acte_name)
            actes_set = sorted(actes_set) or ["—"]

            # Label d'affichage de la préfecture
            if pref_name == "Hors Plateforme":
                pref_label = pref_name
            elif pref_name in PREFECTURES_SET:
                pref_label = f"Préfecture de {pref_name.capitalize()}"
            else:
                pref_label = f"Province de {pref_name.capitalize()}"

            if len(actes_set) == 1:
                row_copy = copy.deepcopy(tmpl_single)
                tcs = row_copy.findall(f'{{{NS_A}}}tc')
                for tc in tcs:
                    tc.attrib.pop('rowSpan', None)
                    tc.attrib.pop('vMerge', None)
                self._set_tc_text(tcs[0], pref_label)
                self._set_tc_text(tcs[1], str(pref_count))
                self._set_tc_text(tcs[2], actes_set[0])
                tbl_elem.append(row_copy)
            else:
                # Première ligne avec rowSpan
                first_copy = copy.deepcopy(tmpl_multi_first)
                tcs = first_copy.findall(f'{{{NS_A}}}tc')
                tcs[0].set('rowSpan', str(len(actes_set)))
                tcs[1].set('rowSpan', str(len(actes_set)))
                tcs[0].attrib.pop('vMerge', None)
                tcs[1].attrib.pop('vMerge', None)
                self._set_tc_text(tcs[0], pref_label)
                self._set_tc_text(tcs[1], str(pref_count))
                self._set_tc_text(tcs[2], actes_set[0])
                tbl_elem.append(first_copy)
                # Lignes de continuation
                for acte in actes_set[1:]:
                    cont_copy = copy.deepcopy(tmpl_multi_cont)
                    tcs = cont_copy.findall(f'{{{NS_A}}}tc')
                    tcs[0].set('vMerge', '1')
                    tcs[1].set('vMerge', '1')
                    tcs[0].attrib.pop('rowSpan', None)
                    tcs[1].attrib.pop('rowSpan', None)
                    self._set_tc_text(tcs[2], acte)
                    tbl_elem.append(cont_copy)

        # Ligne Total
        total_copy = copy.deepcopy(tmpl_total)
        tcs = total_copy.findall(f'{{{NS_A}}}tc')
        if len(tcs) > 1:
            self._set_tc_text(tcs[1], str(total_count))
        tbl_elem.append(total_copy)

    # ------------------------------------------------------------------
    # Remplissage Slide 2
    # ------------------------------------------------------------------

    def _fill_slide2(self, slide, idx, data):
        dossier = data['dossier']
        projet = data['projet']

        for shape in slide.shapes:
            name = shape.name

            if name == 'ZoneTexte 13':
                self._set_shape_run0(shape, f"{idx}) {data['projet_name']}")

            elif name == 'ZoneTexte 1':
                self._set_shape_run0(shape, dossier.reference or '')

            elif name == 'Google Shape;174;p21':
                tbl = shape.table
                self._set_cell_text(tbl.cell(0, 1), data['location_str'])
                self._set_cell_text(tbl.cell(1, 1), data['foncier_str'])
                self._set_cell_text(tbl.cell(4, 1), data['petitionnaire_str'])
                self._set_cell_text(tbl.cell(5, 1), projet.dispositions_urbanistiques if projet else '')
                self._set_cell_multiline(tbl.cell(6, 1), projet.consistance if projet else '')
                mt_str = f"{data['mt_invest_mdh']:.2f}" if data['mt_invest_mdh'] else "0.00"
                self._set_cell_text(tbl.cell(7, 1), mt_str)
                self._set_cell_text(tbl.cell(7, 3), str(projet.nombre_emploi if projet else 0))

            elif name == 'Tableau 15':
                tbl = shape.table
                self._set_cell_text(tbl.cell(0, 1), data['actes_str'])
                self._set_cell_multiline(tbl.cell(1, 1), data['historique_str'])

            elif name == 'Tableau 20':
                tbl = shape.table
                self._set_cell_multiline(tbl.cell(1, 0), dossier.remarque_crui or '')
                self._set_cell_text(
                    tbl.cell(2, 0),
                    f"Proposition de décision: {dossier.decision_crui or ''}"
                )

        # Images Plan de Masse et Plan de Situation
        # Coordonnées issues du template (en EMU) : position table + hauteur header row
        if projet and projet.plan_masse:
            self._add_image_to_slide(
                slide, projet.plan_masse,
                left=9495884, top=7208846, width=5109496, height=3127007
            )
        if projet and projet.plan_situation:
            self._add_image_to_slide(
                slide, projet.plan_situation,
                left=9495882, top=2203717, width=5109496, height=3053920
            )

    # ------------------------------------------------------------------
    # Duplication de slide
    # ------------------------------------------------------------------

    def _duplicate_slide(self, prs, template_slide):
        from pptx.opc.package import _Relationship

        slide_layout = template_slide.slide_layout
        new_slide = prs.slides.add_slide(slide_layout)

        # Remplacer le spTree par une copie profonde du template
        template_sp_tree = template_slide.shapes._spTree
        new_sp_tree = new_slide.shapes._spTree
        for elem in list(new_sp_tree):
            new_sp_tree.remove(elem)
        for elem in template_sp_tree:
            new_sp_tree.append(copy.deepcopy(elem))

        # Copier les relationships en préservant les rId
        new_rels = new_slide.part._rels._rels
        base_uri = new_slide.part._rels._base_uri
        for rId, rel in template_slide.part._rels._rels.items():
            if 'slideLayout' in rel._reltype:
                continue
            if rId not in new_rels:
                try:
                    new_rels[rId] = _Relationship(
                        base_uri=base_uri,
                        rId=rId,
                        reltype=rel._reltype,
                        target_mode=rel._target_mode,
                        target=rel._target,
                    )
                except Exception:
                    pass

        return new_slide

    # ------------------------------------------------------------------
    # Helpers texte
    # ------------------------------------------------------------------

    def _add_image_to_slide(self, slide, image_b64, left, top, width, height):
        """
        Ajoute une image (base64) sur la slide aux coordonnées EMU indiquées.
        L'image est positionnée par-dessus le cadre réservé dans le template.
        """
        try:
            img_data = base64.b64decode(image_b64)
            img_stream = BytesIO(img_data)
            slide.shapes.add_picture(img_stream, left, top, width, height)
        except Exception as e:
            _logger.warning(f"Impossible d'insérer l'image : {e}")

    def _set_rpr_font(self, rpr_elem, font_name, size_pt):
        """Applique la police et la taille sur un élément <a:rPr> lxml."""
        rpr_elem.set('sz', str(size_pt * 100))
        for latin in rpr_elem.findall(f'{{{NS_A}}}latin'):
            rpr_elem.remove(latin)
        latin = etree.SubElement(rpr_elem, f'{{{NS_A}}}latin')
        latin.set('typeface', font_name)

    def _ensure_txbody(self, tc_elem):
        """Crée un txBody minimal dans une cellule vide si absent."""
        txBody = tc_elem.find(f'{{{NS_A}}}txBody')
        if txBody is None:
            txBody = etree.SubElement(tc_elem, f'{{{NS_A}}}txBody')
            etree.SubElement(txBody, f'{{{NS_A}}}bodyPr')
            etree.SubElement(txBody, f'{{{NS_A}}}lstStyle')
        paras = txBody.findall(f'{{{NS_A}}}p')
        if not paras:
            etree.SubElement(txBody, f'{{{NS_A}}}p')
        return txBody

    def _insert_run_before_endpararpr(self, p, new_text, rpr_copy):
        """Insère <a:r> avant <a:endParaRPr> (requis par le spec OOXML)."""
        a_r = etree.Element(f'{{{NS_A}}}r')
        a_r.append(rpr_copy)
        a_t = etree.SubElement(a_r, f'{{{NS_A}}}t')
        a_t.text = str(new_text) if new_text else ''
        end_rpr = p.find(f'{{{NS_A}}}endParaRPr')
        if end_rpr is not None:
            end_rpr.addprevious(a_r)
        else:
            p.append(a_r)

    def _set_tc_text(self, tc_elem, new_text):
        """Remplace le texte d'une cellule <a:tc> brute (lxml) — Calibri 12pt."""
        txBody = self._ensure_txbody(tc_elem)
        paras = txBody.findall(f'{{{NS_A}}}p')
        p = paras[0]
        for r in p.findall(f'{{{NS_A}}}r'):
            p.remove(r)
        for extra_p in paras[1:]:
            txBody.remove(extra_p)
        rpr = etree.Element(f'{{{NS_A}}}rPr')
        rpr.set('lang', 'fr-FR')
        rpr.set('dirty', '0')
        self._set_rpr_font(rpr, 'Calibri', 12)
        self._insert_run_before_endpararpr(p, new_text, rpr)

    def _set_cell_text(self, cell, text):
        """Remplace le texte d'une cellule python-pptx."""
        self._set_tc_text(cell._tc, text)

    def _set_cell_multiline(self, cell, text):
        """Remplace le texte multilignes d'une cellule — Calibri 12pt, une <a:p> par ligne."""
        txBody = self._ensure_txbody(cell._tc)
        paras = txBody.findall(f'{{{NS_A}}}p')

        # Extraire le format du premier run disponible
        rpr_template = None
        for p in paras:
            for r in p.findall(f'{{{NS_A}}}r'):
                rpr = r.find(f'{{{NS_A}}}rPr')
                if rpr is not None:
                    rpr_template = copy.deepcopy(rpr)
                    break
            if rpr_template is not None:
                break
        if rpr_template is None:
            rpr_template = etree.Element(f'{{{NS_A}}}rPr')
            rpr_template.set('lang', 'fr-FR')
            rpr_template.set('dirty', '0')
        self._set_rpr_font(rpr_template, 'Calibri', 12)

        # Supprimer tous les paragraphes existants
        for p in paras:
            txBody.remove(p)

        lines = (text or '').split('\n') if text else ['']
        for line in lines:
            new_p = etree.SubElement(txBody, f'{{{NS_A}}}p')
            epr = etree.SubElement(new_p, f'{{{NS_A}}}endParaRPr')
            epr.set('lang', 'fr-FR')
            epr.set('dirty', '0')
            if line:
                rpr_line = copy.deepcopy(rpr_template)
                self._insert_run_before_endpararpr(new_p, line, rpr_line)

    def _set_shape_run0(self, shape, text):
        """Remplace le texte du premier run — Calibri 26pt (intitulé / référence dossier)."""
        tf = shape.text_frame
        if not tf.paragraphs:
            return
        para = tf.paragraphs[0]
        if not para.runs:
            return
        run = para.runs[0]
        run.text = str(text) if text else ''
        # Appliquer Calibri 26pt via XML directement (1pt = 12700 EMU, sz en centièmes de pt)
        rpr = run._r.find(f'{{{NS_A}}}rPr')
        if rpr is None:
            rpr = etree.Element(f'{{{NS_A}}}rPr')
            rpr.set('lang', 'fr-FR')
            rpr.set('dirty', '0')
            run._r.insert(0, rpr)
        self._set_rpr_font(rpr, 'Calibri', 26)
        for extra in list(para.runs)[1:]:
            para._p.remove(extra._r)
